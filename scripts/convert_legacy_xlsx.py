#!/usr/bin/env python3
"""
Converts the legacy "grid" dock schedule spreadsheet (one tab per year, berths
as rows, calendar days as columns) into normalized JSON the app can seed into
its database: data/berths.json and data/reservations.json.

Why this exists: the source workbook is a manually-maintained visual grid.
Reading it means re-implementing the same "scan the grid" process a human
dock coordinator used to do by eye - this script documents that process
explicitly instead of hiding it, since it's the clearest way to explain how
the historical data was interpreted.

Usage:
    python3 scripts/convert_legacy_xlsx.py

Reads:  data/source/dock-schedule-sample.xlsx
Writes: data/berths.json, data/reservations.json
"""
import calendar
import json
import re
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE_XLSX = ROOT / "data" / "source" / "dock-schedule-sample.xlsx"
BERTHS_OUT = ROOT / "data" / "berths.json"
RESERVATIONS_OUT = ROOT / "data" / "reservations.json"

MONTHS = [
    "JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
    "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER",
]
MONTH_INDEX = {m: i + 1 for i, m in enumerate(MONTHS)}

# Row label pattern used for normal named berths, e.g. "North Pier West - 410'"
BERTH_LABEL_RE = re.compile(r"^(?P<name>.+?)\s*-\s*(?P<len>\d+)\s*'?\s*$")

# Row/column header pattern for month blocks: "AUGUST 1997" or just "August"
# (in which case the sheet's own year applies).
MONTH_HEADER_RE = re.compile(
    r"^(?P<month>" + "|".join(MONTHS) + r")\s*(?P<year>\d{4})?\s*$", re.I
)

# Occupant-name prefixes that identify an entry as a vessel rather than an
# event/operational note. Anything else (e.g. "Community sail day",
# "Bollard replacement, west face", "ETA 1400") is treated as an EVENT.
VESSEL_PREFIXES = ("R/V", "M/V", "OSV", "BARGE", "S/V", "F/V", "M/Y", "TUG", "S/Y")

FINGER_PIER_LABEL = "North Finger Piers"


def classify_occupant(text: str) -> str:
    upper = text.strip().upper()
    for prefix in VESSEL_PREFIXES:
        if upper == prefix or upper.startswith(prefix + " ") or upper.startswith(prefix + "/"):
            return "VESSEL"
    return "EVENT"


def build_day_map(rows, header_row_idx, max_col, year, month):
    """Collect every numeric cell across the header rows following a month
    title (day numbers may be sparse - sometimes only day 1 is filled in -
    or dense - every column labeled). Anchor on whatever numbers exist and
    extrapolate, since columns always represent consecutive calendar days."""
    days_in_month = calendar.monthrange(year, month)[1]
    anchors = {}
    for r in range(header_row_idx, min(header_row_idx + 3, len(rows))):
        row = rows[r]
        first_cell = row[0].value if len(row) else None
        if isinstance(first_cell, str) and BERTH_LABEL_RE.match(first_cell.strip()):
            break
        if isinstance(first_cell, str) and FINGER_PIER_LABEL in first_cell:
            break
        for c in range(1, max_col):
            v = row[c].value if c < len(row) else None
            if isinstance(v, (int, float)) and 1 <= v <= 31:
                anchors[c] = int(v)
    if not anchors:
        return {}
    any_col, any_day = next(iter(anchors.items()))
    day_map = {}
    for c in range(0, max_col):
        day = any_day + (c - any_col)
        if 1 <= day <= days_in_month:
            day_map[c] = day
    return day_map


def parse_sheet(ws, sheet_year):
    rows = list(ws.iter_rows())
    max_col = ws.max_column
    cells = []  # (berth_name, berth_len_or_none, date, occupant_text)

    day_map = {}
    current_year = sheet_year
    current_month = None
    active_berth = None  # (name, length_or_none)
    finger_piers_continuation = 0

    i = 0
    while i < len(rows):
        row = rows[i]
        a = row[0].value

        if isinstance(a, str):
            text = a.strip()
            m = MONTH_HEADER_RE.match(text)
            if m:
                current_month = MONTH_INDEX[m.group("month").upper()]
                current_year = int(m.group("year")) if m.group("year") else sheet_year
                day_map = build_day_map(rows, i, max_col, current_year, current_month)
                active_berth = None
                finger_piers_continuation = 0
                i += 1
                continue

            berth_match = BERTH_LABEL_RE.match(text)
            if berth_match and current_month is not None:
                active_berth = (berth_match.group("name").strip(), int(berth_match.group("len")))
                finger_piers_continuation = 0
                for c, day in day_map.items():
                    v = row[c].value if c < len(row) else None
                    if isinstance(v, str) and v.strip():
                        d = date(current_year, current_month, day)
                        cells.append((active_berth[0], active_berth[1], d, v.strip()))
                i += 1
                continue

            if FINGER_PIER_LABEL in text and current_month is not None:
                active_berth = (FINGER_PIER_LABEL, None)
                finger_piers_continuation = 2  # allows the 2 rows that follow
                for c, day in day_map.items():
                    v = row[c].value if c < len(row) else None
                    if isinstance(v, str) and v.strip():
                        d = date(current_year, current_month, day)
                        cells.append((FINGER_PIER_LABEL, None, d, v.strip()))
                i += 1
                continue

            if "Small craft slips" in text and active_berth and active_berth[0] == FINGER_PIER_LABEL:
                for c, day in day_map.items():
                    v = row[c].value if c < len(row) else None
                    if isinstance(v, str) and v.strip():
                        d = date(current_year, current_month, day)
                        cells.append((FINGER_PIER_LABEL, None, d, v.strip()))
                finger_piers_continuation -= 1
                i += 1
                continue

            # Any other labeled row (facility title, contact info, etc.) - not
            # schedule data. Reset state so we don't misattribute later cells.
            active_berth = None
            finger_piers_continuation = 0
            i += 1
            continue

        # Blank first-column row.
        if active_berth and active_berth[0] == FINGER_PIER_LABEL and finger_piers_continuation > 0:
            for c, day in day_map.items():
                v = row[c].value if c < len(row) else None
                if isinstance(v, str) and v.strip():
                    d = date(current_year, current_month, day)
                    cells.append((FINGER_PIER_LABEL, None, d, v.strip()))
            finger_piers_continuation -= 1
        else:
            active_berth = None
            finger_piers_continuation = 0
        i += 1

    return cells


def merge_runs(cells):
    """Collapse consecutive calendar days with the same occupant string, on
    the same berth, into a single reservation date range. Works across month
    and year-sheet boundaries because everything is keyed on absolute dates."""
    by_berth = {}
    for name, length, d, occupant in cells:
        by_berth.setdefault(name, {"length": length, "entries": []})
        if length is not None:
            by_berth[name]["length"] = length
        by_berth[name]["entries"].append((d, occupant))

    reservations = []
    for berth_name, info in by_berth.items():
        # A few year tabs duplicate the tail end of the prior year's December
        # block at their own start (a copy-paste artifact in the source
        # workbook). Dedupe identical (date, occupant) pairs so that doesn't
        # get read as two separate bookings.
        entries = sorted(set(info["entries"]), key=lambda e: e[0])
        run_start = run_end = run_occupant = None
        for d, occupant in entries:
            if run_occupant == occupant and d == run_end + timedelta(days=1):
                run_end = d
                continue
            if run_occupant is not None:
                reservations.append((berth_name, run_occupant, run_start, run_end))
            run_start = run_end = d
            run_occupant = occupant
        if run_occupant is not None:
            reservations.append((berth_name, run_occupant, run_start, run_end))
    return reservations


def find_overlap_clusters(reservations):
    """For berths with a fixed single-slot capacity, count how many
    reservation pairs on the same berth overlap in date range. This is the
    manual-grid-checking problem the project is meant to replace, quantified
    against the real 23 years of sample data."""
    by_berth = {}
    for berth_name, occupant, start, end in reservations:
        if berth_name == FINGER_PIER_LABEL:
            continue  # multi-slip area, overlap is expected there
        by_berth.setdefault(berth_name, []).append((start, end, occupant))

    conflicts = []
    for berth_name, spans in by_berth.items():
        spans.sort()
        for idx, (s1, e1, occ1) in enumerate(spans):
            for s2, e2, occ2 in spans[idx + 1:]:
                if s2 > e1:
                    break
                if s2 <= e1 and s1 <= e2:
                    conflicts.append((berth_name, occ1, s1, e1, occ2, s2, e2))
    return conflicts


def main():
    wb = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    year_sheets = [n for n in wb.sheetnames if n.isdigit()]

    all_cells = []
    for name in sorted(year_sheets):
        ws = wb[name]
        all_cells.extend(parse_sheet(ws, int(name)))

    reservations = merge_runs(all_cells)

    berth_lengths = {}
    for name, length, _, _ in all_cells:
        if length is not None:
            berth_lengths[name] = length
    if FINGER_PIER_LABEL not in berth_lengths:
        berth_lengths[FINGER_PIER_LABEL] = None

    berths_json = [
        {
            "name": name,
            "lengthFt": length,
            "capacity": None if name == FINGER_PIER_LABEL else 1,
        }
        for name, length in sorted(berth_lengths.items())
    ]

    reservations_json = [
        {
            "berthName": berth_name,
            "occupantName": occupant,
            "occupantType": classify_occupant(occupant),
            "vesselLengthFt": None,
            "startDate": start.isoformat(),
            "endDate": end.isoformat(),
            "notes": "Imported from 23-year legacy schedule spreadsheet.",
        }
        for berth_name, occupant, start, end in sorted(reservations, key=lambda r: (r[0], r[2]))
    ]

    BERTHS_OUT.write_text(json.dumps(berths_json, indent=2) + "\n")
    RESERVATIONS_OUT.write_text(json.dumps(reservations_json, indent=2) + "\n")

    conflicts = find_overlap_clusters(reservations)

    print(f"Parsed {len(year_sheets)} year sheets ({min(year_sheets)}-{max(year_sheets)})")
    print(f"Berths found: {len(berths_json)} -> {[b['name'] for b in berths_json]}")
    print(f"Reservations generated: {len(reservations_json)}")
    print(f"Historical double-booking conflicts detected (single-slot berths): {len(conflicts)}")
    if conflicts:
        print("Sample conflicts:")
        for c in conflicts[:5]:
            berth, occ1, s1, e1, occ2, s2, e2 = c
            print(f"  [{berth}] '{occ1}' {s1}..{e1}  overlaps  '{occ2}' {s2}..{e2}")


if __name__ == "__main__":
    main()
